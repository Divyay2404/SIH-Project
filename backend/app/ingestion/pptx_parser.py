"""
Structure-aware PPTX parser using python-pptx.
Extracts readable slide titles, shape text, and table text into
flat chunk dictionaries compatible with the existing vector store,
RAG queries, and PDF-style page reconstruction.
"""

import os
import re
from typing import Any, Dict, List

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False


class PPTXStructureParser:
    """Parses .pptx lecture decks into slide-scoped chunk dictionaries."""

    def __init__(self):
        self.available = PPTX_AVAILABLE

    def parse_pptx(self, file_path: str) -> List[Dict[str, Any]]:
        """Return flat chunk list with page=slide number and paragraph text."""
        if not self.available:
            raise RuntimeError("python-pptx is required to parse PowerPoint files.")

        if not os.path.exists(file_path):
            raise FileNotFoundError(f"PPTX file not found at: {file_path}")

        prs = Presentation(file_path)
        chunks: List[Dict[str, Any]] = []

        for slide_index, slide in enumerate(prs.slides, start=1):
            slide_lines: List[str] = []
            slide_title = ""

            if slide.shapes.title is not None and slide.shapes.title.text.strip():
                slide_title = slide.shapes.title.text.strip()
                slide_lines.append(slide_title)

            for shape in slide.shapes:
                text_value = ""
                if hasattr(shape, "has_text_frame") and shape.has_text_frame:
                    paragraphs = []
                    for paragraph in shape.text_frame.paragraphs:
                        p_text = paragraph.text.strip()
                        if p_text:
                            paragraphs.append(p_text)
                    if paragraphs:
                        text_value = "\n".join(paragraphs)
                        slide_lines.append(text_value)

                if hasattr(shape, "has_table") and shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        row_text = []
                        for cell in row.cells:
                            cell_text = cell.text.strip()
                            if cell_text:
                                row_text.append(cell_text)
                        if row_text:
                            slide_lines.append(" | ".join(row_text))

            clean_lines = []
            seen = set()
            for line in slide_lines:
                clean_line = re.sub(r"\s+", " ", line).strip()
                if not clean_line:
                    continue
                low = clean_line.lower()
                if low not in seen:
                    seen.add(low)
                    clean_lines.append(clean_line)

            slide_text = "\n".join(clean_lines)
            if not slide_text:
                continue

            chunks.append({
                "page": slide_index,
                "page_index": slide_index - 1,
                "text": slide_text,
                "is_heading": bool(slide_title),
                "bbox": [50.0, 100.0, 500.0, 220.0],
                "native_bbox": [50.0, 100.0, 500.0, 220.0],
                "block_type": "pptx_slide",
            })

        return chunks


pptx_parser_engine = PPTXStructureParser()
